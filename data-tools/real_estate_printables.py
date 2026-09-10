"""Normalize semi-structured real-estate workbooks and generate searchable PDFs.

The program reads vertically stacked account sections and horizontally unpivoted
payment ledgers from legacy Excel worksheets. It writes normalized CSV files and
a SQLite database, and includes a Tkinter UI for name search and PDF generation.
"""

from __future__ import annotations

import argparse
import json
import re
import sqlite3
import sys
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

import pandas as pd
from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


APP_DIR = Path(__file__).resolve().parent
DEFAULT_SOURCE = APP_DIR / "Sample.xlsx"
DEFAULT_OUTPUT_DIR = APP_DIR / "normalized_output"

ACCOUNT_COLUMNS = [
    "account_id", "sheet_name", "source_row", "account_category",
    "reservation_date", "location_or_origin", "primary_client_name",
    "secondary_client_or_note", "commission_pct", "commission_raw", "sqm",
    "sqm_raw", "dicer_agent", "block", "lot", "contact_number",
    "list_price", "down_payment", "stated_balance", "monthly_amortization",
    "monthly_amortization_raw", "payment_terms", "due_date", "due_date_source",
]
PAYMENT_COLUMNS = [
    "payment_id", "account_id", "payment_date", "amount", "payment_type",
    "or_ar_number", "raw_remark",
]
EVENT_COLUMNS = [
    "event_id", "account_id", "event_type", "event_date", "event_amount",
    "target_party_or_project", "raw_text",
]
DIAGNOSTIC_COLUMNS = ["sheet_name", "status", "message"]

HEADER_ALIASES = {
    "reservation_date": {"datepaid", "reservationdate", "date"},
    "primary_client_name": {"name", "client", "vendee"},
    "commission_pct": {"%", "percent", "commission", "commissionpct"},
    "sqm": {"sqm", "area", "areasqm"},
    "dicer_agent": {"dicer", "agent", "diceragent"},
    "block": {"blk", "block"},
    "lot": {"lot", "lotno"},
    "contact_number": {"cpno", "contactno", "contactnumber", "phone"},
    "list_price": {"lp", "lotprice", "listprice"},
    "down_payment": {"dp", "downpayment"},
    "stated_balance": {"bal", "balance", "statedbalance"},
    "monthly_amortization": {"monthly", "amortization", "monthlypayment"},
    "payment_terms": {"terms", "term", "paymentterms"},
}

SECTION_PATTERNS = [
    (re.compile(r"\bBACK\s*OUT\b", re.I), "Back Out"),
    (re.compile(r"\b(FULL\s*REFUND|ASSUMED\s*BY)\b", re.I), "Full Refund/Assumed"),
    (re.compile(r"\bTRANSFER(?:RED)?\b", re.I), "Transferred"),
    (re.compile(r"\bFORFEITED?\b", re.I), "Forfeited"),
]
FEE_PATTERNS = [
    (re.compile(r"\bDAS\b", re.I), "DAS Fee"),
    (re.compile(r"\bDCS\b", re.I), "DCS Fee"),
    (re.compile(r"\bSPA\b", re.I), "SPA Fee"),
    (re.compile(r"\bTITL(?:ING|E)?\b", re.I), "Titling Fee"),
]
DATE_TOKEN_RE = re.compile(
    r"\b(?:\d{4}[-/]\d{1,2}[-/]\d{1,2}|\d{1,2}[-/]\d{1,2}[-/]\d{2,4}|\d{1,2}-[A-Za-z]{3}-\d{2,4})\b"
)
MONEY_TOKEN_RE = re.compile(r"(?<![\d/])(?:\d{1,3}(?:,\d{3})+|\d+(?:\.\d+)?)(?![\d/])")


def normalize_text(value: Any) -> str | None:
    """Return compact text or None for blank / NaN values."""
    if value is None or value is pd.NA or (isinstance(value, float) and pd.isna(value)):
        return None
    text = re.sub(r"\s+", " ", str(value).replace("\u00a0", " ")).strip(" ,\t\r\n")
    return text or None


def canonical_header(value: Any) -> str:
    text = normalize_text(value) or ""
    if text == "%":
        return "%"
    return re.sub(r"[^a-z0-9]", "", text.lower())


def parse_number(value: Any) -> float | None:
    """Parse ordinary, comma-separated, scientific, or annotated numeric values."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float, Decimal)):
        return float(value)

    text = normalize_text(value)
    if not text or text.lower() in {"-", "n/a", "na", "none"}:
        return None
    negative = bool(re.fullmatch(r"\([^)]*\)", text))
    cleaned = text.replace("₱", "").replace("P", "").replace(" ", "")
    match = re.search(r"[-+]?\d+(?:,\d{3})*(?:\.\d+)?(?:[eE][-+]?\d+)?", cleaned)
    if not match:
        return None
    try:
        result = float(Decimal(match.group(0).replace(",", "")))
        return -abs(result) if negative else result
    except (InvalidOperation, ValueError):
        return None


def parse_date(value: Any) -> date | None:
    """Parse Excel dates and explicit date text without treating place names as dates."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and 1000 < float(value) < 60000:
        return (datetime(1899, 12, 30) + timedelta(days=float(value))).date()
    text = normalize_text(value)
    # Avoid a costly date-parser call for ledger labels such as "DAS", names,
    # and project names.  This also prevents a location in column A becoming a
    # spurious reservation date.
    if not text or not DATE_TOKEN_RE.fullmatch(text):
        return None
    try:
        parsed = pd.to_datetime(text, errors="raise", dayfirst=False)
        if pd.isna(parsed):
            return None
        return parsed.date()
    except (ValueError, TypeError, OverflowError):
        return None


def parse_date_from_text(text: str) -> date | None:
    """Extract the first explicit date token from an inline ledger remark."""
    match = DATE_TOKEN_RE.search(text)
    return parse_date(match.group(0)) if match else None


def normalize_contact(value: Any) -> str | None:
    """Return one canonical 10- or 11-digit Philippine-style contact number."""
    if value is None:
        return None
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        try:
            text = format(Decimal(str(value)), "f").split(".")[0]
        except InvalidOperation:
            text = str(value)
    else:
        text = normalize_text(value) or ""
    candidates = re.findall(r"\d{9,14}", re.sub(r"[^0-9/]", "", text))
    if not candidates:
        return None
    digits = candidates[0]
    if digits.startswith("63") and len(digits) >= 12:
        digits = "0" + digits[2:]
    elif len(digits) == 10 and digits.startswith("9"):
        digits = "0" + digits
    if len(digits) in {10, 11}:
        return digits
    return digits[:11] if len(digits) > 11 else None


def parse_name(value: Any) -> tuple[str | None, str | None]:
    """Split a primary client name from parenthetical or slash-separated secondary names."""
    text = normalize_text(value)
    if not text:
        return None, None
    secondary_parts: list[str] = []
    parenthetical = re.findall(r"\(([^)]+)\)", text)
    if parenthetical:
        secondary_parts.extend(normalize_text(part) or "" for part in parenthetical)
        text = re.sub(r"\([^)]*\)", "", text).strip(" ,")
    if "/" in text:
        split = [normalize_text(part) for part in text.split("/")]
        text = split[0] or ""
        secondary_parts.extend(part for part in split[1:] if part)
    return normalize_text(text), normalize_text("; ".join(part for part in secondary_parts if part))


def parse_commission(value: Any) -> tuple[float | None, str | None]:
    raw = normalize_text(value)
    return parse_number(raw), raw


def parse_sqm(value: Any) -> tuple[float | None, str | None]:
    raw = normalize_text(value)
    return parse_number(raw), raw


def derive_due_date(reservation_date: date | None, terms: Any) -> tuple[date | None, str | None]:
    """Derive due date only when the term text makes the rule explicit."""
    term_text = normalize_text(terms)
    if not reservation_date or not term_text:
        return None, None
    upper = term_text.upper()
    if "CASH" in upper or upper in {"0", "0.0"}:
        return reservation_date, "reservation_date_for_cash"
    year_match = re.search(r"(\d+(?:\.\d+)?)\s*(?:YRS?|YEARS?)\b", upper)
    if year_match:
        years = int(float(year_match.group(1)))
        try:
            return reservation_date.replace(year=reservation_date.year + years), "reservation_date_plus_terms_years"
        except ValueError:  # Feb 29 -> Feb 28 in a non-leap target year
            return reservation_date.replace(month=2, day=28, year=reservation_date.year + years), "reservation_date_plus_terms_years"
    month_match = re.search(r"(\d+)\s*(?:MOS?|MONTHS?)\b", upper)
    if month_match:
        months = int(month_match.group(1))
        year = reservation_date.year + (reservation_date.month - 1 + months) // 12
        month = (reservation_date.month - 1 + months) % 12 + 1
        day = min(reservation_date.day, [31, 29 if year % 4 == 0 and (year % 100 != 0 or year % 400 == 0) else 28,
                                         31, 30, 31, 30, 31, 31, 30, 31, 30, 31][month - 1])
        return date(year, month, day), "reservation_date_plus_terms_months"
    return None, None


def detect_category(row_values: Iterable[Any]) -> str | None:
    text = " | ".join(normalize_text(value) or "" for value in row_values)
    for pattern, category in SECTION_PATTERNS:
        if pattern.search(text):
            return category
    return None


def find_header_row(ws: Any) -> tuple[int, dict[str, int]]:
    """Find a row containing the majority of the expected account headers."""
    best: tuple[int, dict[str, int]] | None = None
    for row_index in range(1, min(ws.max_row, 30) + 1):
        mapping: dict[str, int] = {}
        for column_index in range(1, ws.max_column + 1):
            header = canonical_header(ws.cell(row_index, column_index).value)
            for field, aliases in HEADER_ALIASES.items():
                if header in aliases and field not in mapping:
                    mapping[field] = column_index - 1
        if best is None or len(mapping) > len(best[1]):
            best = (row_index, mapping)
    if not best or len(best[1]) < 8:
        raise ValueError(f"Could not identify the account header row in sheet '{ws.title}'.")
    return best


def is_summary_row(row_data: dict[str, Any]) -> bool:
    name = normalize_text(row_data.get("primary_client_name"))
    values = [row_data.get("list_price"), row_data.get("down_payment"), row_data.get("stated_balance")]
    row_text = " ".join(normalize_text(value) or "" for value in row_data.values())
    return bool(re.search(r"\b(total|subtotal|grand total)\b", row_text, re.I)) or (
        not name and any(parse_number(value) is not None for value in values)
    )


def load_raw_sheet_with_sections(
    file_path: str | Path, sheet_name: str, workbook: Any | None = None
) -> list[dict[str, Any]]:
    """Read one worksheet, locate sections, and return raw account rows plus ledger cells."""
    # Normal mode makes random cell access fast.  Read-only mode reparses the
    # sheet on every ``cell(row, column)`` call and is prohibitively slow here.
    owns_workbook = workbook is None
    if owns_workbook:
        workbook = load_workbook(file_path, read_only=False, data_only=True)
    try:
        worksheet = workbook[sheet_name]
        header_row, header_map = find_header_row(worksheet)
        ledger_start = max(header_map.values()) + 1
        category = "Active"
        records: list[dict[str, Any]] = []

        for row_index in range(header_row + 1, worksheet.max_row + 1):
            raw_values = [worksheet.cell(row_index, column).value for column in range(1, worksheet.max_column + 1)]
            if not any(value is not None for value in raw_values):
                continue
            row_data = {field: raw_values[index] for field, index in header_map.items()}
            section = detect_category(raw_values)
            # A genuine account can contain an "Assumed by" or transfer note in
            # its ledger.  Only a nameless row may switch the current section.
            if section and not normalize_text(row_data.get("primary_client_name")):
                category = section
                continue
            row_data.update({
                "sheet_name": sheet_name,
                "source_row": row_index,
                "account_category": category,
                "ledger_cells": raw_values[ledger_start:],
            })
            if is_summary_row(row_data):
                continue
            if not normalize_text(row_data.get("primary_client_name")):
                continue
            records.append(row_data)
        return records
    finally:
        if owns_workbook:
            workbook.close()


def parse_account_metadata(row_data: dict[str, Any], category: str) -> dict[str, Any]:
    """Normalize one raw account row into the dim_accounts contract."""
    raw_first_column = row_data.get("reservation_date")
    reservation_date = parse_date(raw_first_column)
    location_or_origin = None if reservation_date else normalize_text(raw_first_column)
    primary_name, secondary_name = parse_name(row_data.get("primary_client_name"))
    commission_pct, commission_raw = parse_commission(row_data.get("commission_pct"))
    sqm, sqm_raw = parse_sqm(row_data.get("sqm"))
    monthly_raw = normalize_text(row_data.get("monthly_amortization"))
    monthly_value = parse_number(monthly_raw)
    terms = normalize_text(row_data.get("payment_terms"))
    due_date, due_date_source = derive_due_date(reservation_date, terms)
    account_id = f"{row_data['sheet_name']}_{row_data['source_row']}"

    return {
        "account_id": account_id,
        "sheet_name": row_data["sheet_name"],
        "source_row": row_data["source_row"],
        "account_category": category,
        "reservation_date": reservation_date,
        "location_or_origin": location_or_origin,
        "primary_client_name": primary_name,
        "secondary_client_or_note": secondary_name,
        "commission_pct": commission_pct,
        "commission_raw": commission_raw,
        "sqm": sqm,
        "sqm_raw": sqm_raw,
        "dicer_agent": normalize_text(row_data.get("dicer_agent")),
        "block": normalize_text(row_data.get("block")),
        "lot": normalize_text(row_data.get("lot")),
        "contact_number": normalize_contact(row_data.get("contact_number")),
        "list_price": parse_number(row_data.get("list_price")),
        "down_payment": parse_number(row_data.get("down_payment")),
        "stated_balance": parse_number(row_data.get("stated_balance")),
        "monthly_amortization": monthly_value if monthly_value is not None else monthly_raw,
        "monthly_amortization_raw": monthly_raw,
        "payment_terms": terms,
        "due_date": due_date,
        "due_date_source": due_date_source,
    }


def fee_type_for(text: str) -> str | None:
    for pattern, payment_type in FEE_PATTERNS:
        if pattern.search(text):
            return payment_type
    return None


def extract_refund(text: str) -> tuple[date | None, float | None]:
    without_dates = DATE_TOKEN_RE.sub(" ", text)
    values = [parse_number(token) for token in MONEY_TOKEN_RE.findall(without_dates)]
    values = [value for value in values if value is not None]
    return parse_date_from_text(text), max(values) if values else None


def extract_event(text: str) -> dict[str, Any] | None:
    """Identify explicit refund, assumption, transfer, and forfeiture ledger notes."""
    normalized = normalize_text(text)
    if not normalized:
        return None
    if re.search(r"\bREFUND(?:ED)?\b", normalized, re.I):
        event_date, event_amount = extract_refund(normalized)
        return {
            "event_type": "Refund",
            "event_date": event_date,
            "event_amount": event_amount,
            "target_party_or_project": None,
            "raw_text": normalized,
        }
    assumed = re.search(r"\bASSUMED\s+BY\s*[:;]?\s*(.+)$", normalized, re.I)
    if assumed:
        return {
            "event_type": "Assumed By", "event_date": parse_date_from_text(normalized),
            "event_amount": None, "target_party_or_project": normalize_text(assumed.group(1)),
            "raw_text": normalized,
        }
    transferred = re.search(r"\bTRANSFER(?:\s+OF\s+OWNERSHIP)?\s+TO\s*[:;]?\s*(.+)$", normalized, re.I)
    if not transferred:
        transferred = re.search(r"^\s*TO\s+(.+)$", normalized, re.I)
    if transferred:
        return {
            "event_type": "Transferred To", "event_date": parse_date_from_text(normalized),
            "event_amount": None, "target_party_or_project": normalize_text(transferred.group(1)),
            "raw_text": normalized,
        }
    if re.search(r"\bFORFEITED?\b", normalized, re.I):
        return {
            "event_type": "Forfeited", "event_date": parse_date_from_text(normalized),
            "event_amount": None, "target_party_or_project": None, "raw_text": normalized,
        }
    return None


def next_non_blank(cells: list[Any], start_index: int) -> tuple[int | None, Any | None]:
    for index in range(start_index, len(cells)):
        if normalize_text(cells[index]) is not None:
            return index, cells[index]
    return None, None


def unpivot_payment_ledger(account_id: str, ledger_cells: Iterable[Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Convert horizontal date/amount and fee cells into payments and account events."""
    cells = list(ledger_cells)
    payments: list[dict[str, Any]] = []
    events: list[dict[str, Any]] = []
    index = 0

    while index < len(cells):
        current = cells[index]
        text = normalize_text(current)
        if not text:
            index += 1
            continue

        event = extract_event(text)
        if event:
            event["account_id"] = account_id
            events.append(event)
            if event["event_type"] == "Refund" and event["event_amount"] is not None:
                payments.append({
                    "account_id": account_id,
                    "payment_date": event["event_date"],
                    "amount": -abs(event["event_amount"]),
                    "payment_type": "Refund",
                    "or_ar_number": None,
                    "raw_remark": event["raw_text"],
                })
            index += 1
            continue

        fee_type = fee_type_for(text)
        if fee_type:
            amount_index, amount_cell = next_non_blank(cells, index + 1)
            amount = parse_number(amount_cell)
            if amount is not None:
                payments.append({
                    "account_id": account_id, "payment_date": None, "amount": amount,
                    "payment_type": fee_type, "or_ar_number": None, "raw_remark": text,
                })
                index = (amount_index or index) + 1
                continue
            index += 1
            continue

        payment_date = parse_date(current)
        if payment_date:
            amount_index, amount_cell = next_non_blank(cells, index + 1)
            amount = parse_number(amount_cell)
            if amount is not None:
                payments.append({
                    "account_id": account_id, "payment_date": payment_date, "amount": amount,
                    "payment_type": "Installment", "or_ar_number": None, "raw_remark": None,
                })
                index = (amount_index or index) + 1
                continue
            index += 1
            continue

        # Non-standard labels such as a person's name followed by a payment amount.
        amount_index, amount_cell = next_non_blank(cells, index + 1)
        amount = parse_number(amount_cell)
        if amount is not None:
            payments.append({
                "account_id": account_id, "payment_date": None, "amount": amount,
                "payment_type": "Installment", "or_ar_number": None, "raw_remark": text,
            })
            index = (amount_index or index) + 1
            continue

        # Preserve otherwise unpaired numeric cells for auditability.
        numeric = parse_number(current)
        if numeric is not None:
            payments.append({
                "account_id": account_id, "payment_date": None, "amount": numeric,
                "payment_type": "Installment", "or_ar_number": None,
                "raw_remark": "Unpaired ledger amount",
            })
        index += 1
    return payments, events


def empty_frame(columns: list[str]) -> pd.DataFrame:
    return pd.DataFrame(columns=columns)


def add_category_events(account: dict[str, Any]) -> list[dict[str, Any]]:
    if account["account_category"] != "Forfeited":
        return []
    return [{
        "account_id": account["account_id"], "event_type": "Forfeited", "event_date": None,
        "event_amount": None, "target_party_or_project": None,
        "raw_text": "Account is in the workbook's FORFEITED ACCOUNTS section.",
    }]


def process_entire_workbook(file_path: str | Path) -> dict[str, pd.DataFrame]:
    """Normalize all sheets in a workbook into account, payment, event, and audit frames."""
    file_path = Path(file_path)
    if not file_path.exists():
        raise FileNotFoundError(f"Workbook not found: {file_path}")

    workbook = load_workbook(file_path, read_only=False, data_only=True)
    account_rows: list[dict[str, Any]] = []
    payment_rows: list[dict[str, Any]] = []
    event_rows: list[dict[str, Any]] = []
    diagnostic_rows: list[dict[str, str]] = []

    try:
        for sheet_name in workbook.sheetnames:
            try:
                raw_rows = load_raw_sheet_with_sections(file_path, sheet_name, workbook)
            except ValueError as exc:
                # Some legacy workbooks include title/notes-only tabs.  They are
                # preserved in the diagnostic export instead of stopping the run.
                diagnostic_rows.append({"sheet_name": sheet_name, "status": "Skipped", "message": str(exc)})
                continue
            diagnostic_rows.append({"sheet_name": sheet_name, "status": "Processed", "message": f"{len(raw_rows)} account row(s) found"})
            for raw_row in raw_rows:
                account = parse_account_metadata(raw_row, raw_row["account_category"])
                account_rows.append(account)
                payments, events = unpivot_payment_ledger(account["account_id"], raw_row["ledger_cells"])
                payment_rows.extend(payments)
                event_rows.extend(events)
                event_rows.extend(add_category_events(account))
    finally:
        workbook.close()

    accounts = pd.DataFrame(account_rows, columns=ACCOUNT_COLUMNS)
    payments = pd.DataFrame(payment_rows)
    events = pd.DataFrame(event_rows)
    if payments.empty:
        payments = empty_frame(PAYMENT_COLUMNS[1:])
    if events.empty:
        events = empty_frame(EVENT_COLUMNS[1:])
    payments.insert(0, "payment_id", range(1, len(payments) + 1))
    events.insert(0, "event_id", range(1, len(events) + 1))
    payments = payments.reindex(columns=PAYMENT_COLUMNS)
    events = events.reindex(columns=EVENT_COLUMNS)

    verification = build_balance_verification(accounts, payments)
    return {
        "dim_accounts": accounts,
        "fact_payments": payments,
        "fact_account_events": events,
        "account_balance_verification": verification,
        "ingestion_diagnostics": pd.DataFrame(diagnostic_rows, columns=DIAGNOSTIC_COLUMNS),
    }


def build_balance_verification(accounts: pd.DataFrame, payments: pd.DataFrame) -> pd.DataFrame:
    """Compare stated balance with LP - DP - installment payments, as requested."""
    installments = payments.loc[payments["payment_type"].eq("Installment")].copy()
    paid = installments.groupby("account_id", dropna=False)["amount"].sum().rename("installment_total")
    verification = accounts[["account_id", "primary_client_name", "list_price", "down_payment", "stated_balance"]].copy()
    verification = verification.join(paid, on="account_id")
    verification["installment_total"] = verification["installment_total"].fillna(0.0)
    verification["list_price"] = pd.to_numeric(verification["list_price"], errors="coerce")
    verification["down_payment"] = pd.to_numeric(verification["down_payment"], errors="coerce").fillna(0.0)
    verification["calculated_balance"] = verification["list_price"] - verification["down_payment"] - verification["installment_total"]
    verification["balance_discrepancy"] = verification["stated_balance"] - verification["calculated_balance"]
    verification["status"] = verification["balance_discrepancy"].map(
        lambda value: "Unavailable" if pd.isna(value) else ("Match" if abs(value) < 0.01 else "Review")
    )
    return verification


def serializable_frame(frame: pd.DataFrame) -> pd.DataFrame:
    """Convert Python date values to ISO text for SQLite and CSV stability."""
    result = frame.copy()
    for column in result.columns:
        if result[column].map(lambda value: isinstance(value, (date, datetime))).any():
            result[column] = result[column].map(lambda value: value.isoformat() if isinstance(value, (date, datetime)) else value)
    return result


def export_normalized_data(tables: dict[str, pd.DataFrame], output_dir: str | Path) -> dict[str, Path]:
    """Write normalized CSVs, a SQLite relational database, and a run summary."""
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    paths: dict[str, Path] = {}
    for name, frame in tables.items():
        csv_path = output_dir / f"{name}.csv"
        serializable_frame(frame).to_csv(csv_path, index=False, encoding="utf-8-sig")
        paths[name] = csv_path

    database_path = output_dir / "real_estate_tracker.db"
    with sqlite3.connect(database_path) as connection:
        connection.execute("PRAGMA foreign_keys = ON")
        connection.executescript("""
            DROP TABLE IF EXISTS fact_account_events;
            DROP TABLE IF EXISTS fact_payments;
            DROP TABLE IF EXISTS dim_accounts;
            CREATE TABLE dim_accounts (
                account_id TEXT PRIMARY KEY, sheet_name TEXT NOT NULL, source_row INTEGER NOT NULL,
                account_category TEXT NOT NULL, reservation_date TEXT, location_or_origin TEXT,
                primary_client_name TEXT NOT NULL, secondary_client_or_note TEXT, commission_pct REAL,
                commission_raw TEXT, sqm REAL, sqm_raw TEXT, dicer_agent TEXT, block TEXT, lot TEXT,
                contact_number TEXT, list_price REAL, down_payment REAL, stated_balance REAL,
                monthly_amortization TEXT, monthly_amortization_raw TEXT, payment_terms TEXT,
                due_date TEXT, due_date_source TEXT
            );
            CREATE TABLE fact_payments (
                payment_id INTEGER PRIMARY KEY, account_id TEXT NOT NULL, payment_date TEXT,
                amount REAL NOT NULL, payment_type TEXT NOT NULL, or_ar_number TEXT, raw_remark TEXT,
                FOREIGN KEY(account_id) REFERENCES dim_accounts(account_id)
            );
            CREATE TABLE fact_account_events (
                event_id INTEGER PRIMARY KEY, account_id TEXT NOT NULL, event_type TEXT NOT NULL,
                event_date TEXT, event_amount REAL, target_party_or_project TEXT, raw_text TEXT,
                FOREIGN KEY(account_id) REFERENCES dim_accounts(account_id)
            );
        """)
        serializable_frame(tables["dim_accounts"]).to_sql("dim_accounts", connection, if_exists="append", index=False)
        serializable_frame(tables["fact_payments"]).to_sql("fact_payments", connection, if_exists="append", index=False)
        serializable_frame(tables["fact_account_events"]).to_sql("fact_account_events", connection, if_exists="append", index=False)
    paths["sqlite"] = database_path

    summary = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "accounts": len(tables["dim_accounts"]),
        "payments": len(tables["fact_payments"]),
        "events": len(tables["fact_account_events"]),
        "balance_rows_for_review": int((tables["account_balance_verification"]["status"] == "Review").sum()),
    }
    summary_path = output_dir / "run_summary.json"
    summary_path.write_text(json.dumps(summary, indent=2), encoding="utf-8")
    paths["summary"] = summary_path
    return paths


def money(value: Any) -> str:
    number = parse_number(value)
    return "-" if number is None else f"{number:,.2f}"


def printable_value(value: Any) -> str:
    if value is None or value is pd.NA or (isinstance(value, float) and pd.isna(value)):
        return "-"
    if isinstance(value, (date, datetime)):
        return value.strftime("%B %d, %Y")
    return normalize_text(value) or "-"


def create_printable_pdf(
    account: pd.Series | dict[str, Any],
    payments: pd.DataFrame,
    verification: pd.DataFrame,
    output_dir: str | Path,
) -> Path:
    """Create a static SOA-style PDF for one searched account."""
    account = dict(account)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    safe_name = re.sub(r"[^A-Za-z0-9_-]+", "_", account.get("primary_client_name") or "account").strip("_")
    pdf_path = output_dir / f"{safe_name}_{account['account_id']}.pdf"

    account_payments = payments.loc[payments["account_id"].eq(account["account_id"])].copy()
    account_payments["_sort_date"] = pd.to_datetime(account_payments["payment_date"], errors="coerce")
    account_payments = account_payments.sort_values(["_sort_date", "payment_id"], na_position="last")
    balance_row = verification.loc[verification["account_id"].eq(account["account_id"])].iloc[0]

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("Title", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=18, leading=22,
                                 alignment=TA_CENTER, textColor=colors.HexColor("#12355B"), spaceAfter=4)
    subtitle_style = ParagraphStyle("Subtitle", parent=styles["Normal"], fontSize=9, leading=12, alignment=TA_CENTER,
                                    textColor=colors.HexColor("#4B5563"), spaceAfter=12)
    label_style = ParagraphStyle("Label", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=8, leading=11,
                                 textColor=colors.HexColor("#1F2937"))
    value_style = ParagraphStyle("Value", parent=styles["Normal"], fontSize=9, leading=11, textColor=colors.HexColor("#111827"))
    note_style = ParagraphStyle("Note", parent=styles["Normal"], fontSize=7.5, leading=10, textColor=colors.HexColor("#4B5563"))
    payment_header_style = ParagraphStyle(
        "PaymentHeader", parent=value_style, fontName="Helvetica-Bold", textColor=colors.white
    )

    document = SimpleDocTemplate(str(pdf_path), pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm,
                                 topMargin=14 * mm, bottomMargin=14 * mm)
    story: list[Any] = [
        Paragraph("STATEMENT OF ACCOUNT", title_style),
        Paragraph(f"Date Issued: {datetime.now().strftime('%B %d, %Y')}", subtitle_style),
    ]

    location = normalize_text(account.get("location_or_origin")) or printable_value(account.get("sheet_name"))
    due = account.get("due_date")
    details = [
        ("Vendee Name", printable_value(account.get("primary_client_name"))),
        ("Area Location", printable_value(location)),
        ("Area Size", f"{money(account.get('sqm'))} sq.m." if parse_number(account.get("sqm")) is not None else printable_value(account.get("sqm_raw"))),
        ("Lot Price", money(account.get("list_price"))),
        ("Monthly Payment", money(account.get("monthly_amortization"))),
        ("Terms", printable_value(account.get("payment_terms"))),
        ("Due Date", printable_value(due)),
        ("Account Category", printable_value(account.get("account_category"))),
    ]
    detail_rows = []
    for left, right in zip(details[::2], details[1::2]):
        detail_rows.append([
            Paragraph(left[0], label_style), Paragraph(left[1], value_style),
            Paragraph(right[0], label_style), Paragraph(right[1], value_style),
        ])
    details_table = Table(detail_rows, colWidths=[31 * mm, 58 * mm, 31 * mm, 58 * mm])
    details_table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.extend([details_table, Spacer(1, 8 * mm)])

    payment_rows: list[list[Any]] = [["Date of Payment", "OR/AR #", "Payment Release", "Type / Note"]]
    if account_payments.empty:
        payment_rows.append(["No ledger payments recorded", "-", "-", "-"])
    else:
        for _, payment in account_payments.iterrows():
            remark = normalize_text(payment.get("raw_remark"))
            payment_rows.append([
                printable_value(payment.get("payment_date")),
                printable_value(payment.get("or_ar_number")),
                money(payment.get("amount")),
                remark or printable_value(payment.get("payment_type")),
            ])
    # Paragraph cells wrap long legacy notes rather than overflowing the page.
    payment_rows = [
        [Paragraph(printable_value(cell), payment_header_style if row_number == 0 else value_style) for cell in row]
        for row_number, row in enumerate(payment_rows)
    ]
    payments_table = Table(payment_rows, repeatRows=1, colWidths=[37 * mm, 28 * mm, 38 * mm, 75 * mm])
    payments_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#12355B")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ALIGN", (2, 1), (2, -1), "RIGHT"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.extend([Paragraph("Payment History", label_style), Spacer(1, 2 * mm), payments_table, Spacer(1, 8 * mm)])

    totals = [
        [Paragraph("Payment Release Total", label_style), Paragraph(money(balance_row["installment_total"]), value_style)],
        [Paragraph("Calculated Balance", label_style), Paragraph(money(balance_row["calculated_balance"]), value_style)],
        [Paragraph("Stated Balance", label_style), Paragraph(money(balance_row["stated_balance"]), value_style)],
    ]
    totals_table = Table(totals, colWidths=[64 * mm, 45 * mm], hAlign="RIGHT")
    totals_table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#94A3B8")),
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#EFF6FF")),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(totals_table)
    if balance_row["status"] == "Review":
        story.extend([Spacer(1, 3 * mm), Paragraph(
            f"Reconciliation note: stated and calculated balances differ by {money(balance_row['balance_discrepancy'])}. "
            "Review the source ledger before treating this printable as final.", note_style)])
    due_date_source = normalize_text(account.get("due_date_source"))
    if due_date_source:
        story.extend([Spacer(1, 2 * mm), Paragraph(
            f"Due date source: {due_date_source}. OR/AR values are blank when the workbook does not record them.", note_style)])

    document.build(story)
    return pdf_path


def find_accounts(accounts: pd.DataFrame, query: str) -> pd.DataFrame:
    query = normalize_text(query) or ""
    if not query:
        return accounts.copy()
    mask = accounts["primary_client_name"].fillna("").str.contains(re.escape(query), case=False, regex=True)
    secondary = accounts["secondary_client_or_note"].fillna("").str.contains(re.escape(query), case=False, regex=True)
    return accounts.loc[mask | secondary].copy()


class RealEstateTrackerApp:
    """Small local UI for refresh, name search, and printable generation."""

    def __init__(self, root: Any) -> None:
        import tkinter as tk
        from tkinter import ttk

        self.tk = tk
        self.ttk = ttk
        self.root = root
        self.tables: dict[str, pd.DataFrame] | None = None
        self.source_var = tk.StringVar(value=str(DEFAULT_SOURCE))
        self.search_var = tk.StringVar()
        self.status_var = tk.StringVar(value="Choose Normalize and Export to load the workbook.")
        root.title("Real Estate Printable Tracker")
        root.geometry("1120x650")
        root.minsize(900, 500)
        self._build_layout()

    def _build_layout(self) -> None:
        from tkinter import filedialog, messagebox

        top = self.ttk.Frame(self.root, padding=10)
        top.pack(fill="x")
        self.ttk.Label(top, text="Workbook:").grid(row=0, column=0, sticky="w")
        self.ttk.Entry(top, textvariable=self.source_var, width=100).grid(row=0, column=1, sticky="ew", padx=6)
        top.columnconfigure(1, weight=1)
        self.ttk.Button(top, text="Browse", command=lambda: self._browse(filedialog)).grid(row=0, column=2)
        self.ttk.Button(top, text="Normalize and Export", command=lambda: self._process(messagebox)).grid(row=0, column=3, padx=(6, 0))

        search = self.ttk.Frame(self.root, padding=(10, 0, 10, 8))
        search.pack(fill="x")
        self.ttk.Label(search, text="Search Vendee Name:").pack(side="left")
        entry = self.ttk.Entry(search, textvariable=self.search_var, width=42)
        entry.pack(side="left", padx=6)
        entry.bind("<KeyRelease>", lambda _event: self._refresh_results())
        self.ttk.Button(search, text="Generate PDF for Selected Account", command=lambda: self._generate_pdf(messagebox)).pack(side="right")

        columns = ("account_id", "name", "category", "location", "lot", "balance")
        self.tree = self.ttk.Treeview(self.root, columns=columns, show="headings", selectmode="browse")
        for column, title, width in [
            ("account_id", "Account ID", 125), ("name", "Vendee Name", 270),
            ("category", "Category", 150), ("location", "Area Location", 165),
            ("lot", "Lot", 80), ("balance", "Calculated Balance", 145),
        ]:
            self.tree.heading(column, text=title)
            self.tree.column(column, width=width, anchor="w" if column not in {"balance"} else "e")
        self.tree.pack(fill="both", expand=True, padx=10, pady=(0, 6))
        self.ttk.Label(self.root, textvariable=self.status_var, padding=(10, 2, 10, 10)).pack(fill="x")

    def _browse(self, filedialog: Any) -> None:
        selected = filedialog.askopenfilename(title="Select real-estate workbook", filetypes=[("Excel workbook", "*.xlsx")])
        if selected:
            self.source_var.set(selected)

    def _process(self, messagebox: Any) -> None:
        try:
            self.tables = process_entire_workbook(self.source_var.get())
            paths = export_normalized_data(self.tables, DEFAULT_OUTPUT_DIR)
            self.status_var.set(
                f"Loaded {len(self.tables['dim_accounts'])} accounts, {len(self.tables['fact_payments'])} payments. "
                f"Exports: {paths['sqlite']}"
            )
            self._refresh_results()
        except Exception as exc:
            messagebox.showerror("Normalization failed", str(exc))

    def _refresh_results(self) -> None:
        for item in self.tree.get_children():
            self.tree.delete(item)
        if not self.tables:
            return
        matches = find_accounts(self.tables["dim_accounts"], self.search_var.get())
        verification = self.tables["account_balance_verification"].set_index("account_id")
        for _, row in matches.iterrows():
            calculated = verification.loc[row["account_id"], "calculated_balance"]
            self.tree.insert("", "end", iid=row["account_id"], values=(
                row["account_id"], row["primary_client_name"], row["account_category"],
                normalize_text(row["location_or_origin"]) or printable_value(row["sheet_name"]), row["lot"], money(calculated),
            ))

    def _generate_pdf(self, messagebox: Any) -> None:
        if not self.tables:
            messagebox.showwarning("Load data first", "Select Normalize and Export before generating a printable.")
            return
        selected = self.tree.selection()
        if not selected:
            messagebox.showwarning("Select an account", "Select one search result first.")
            return
        account_id = selected[0]
        account = self.tables["dim_accounts"].loc[self.tables["dim_accounts"]["account_id"].eq(account_id)].iloc[0]
        try:
            pdf_path = create_printable_pdf(
                account, self.tables["fact_payments"], self.tables["account_balance_verification"],
                DEFAULT_OUTPUT_DIR / "printables",
            )
            messagebox.showinfo("Printable created", f"PDF created:\n{pdf_path}")
        except Exception as exc:
            messagebox.showerror("PDF generation failed", str(exc))


def run_ui() -> None:
    import tkinter as tk
    root = tk.Tk()
    RealEstateTrackerApp(root)
    root.mainloop()


def main() -> None:
    parser = argparse.ArgumentParser(description="Normalize real-estate Excel data and generate printables.")
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE, help="Path to the source .xlsx workbook.")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT_DIR, help="Export directory for CSVs, SQLite, and PDFs.")
    parser.add_argument("--build", action="store_true", help="Normalize workbook and export CSV / SQLite files, then exit.")
    parser.add_argument("--pdf", metavar="NAME", help="Build data and create a PDF for the first matching vendee name.")
    args = parser.parse_args()

    if args.build or args.pdf:
        tables = process_entire_workbook(args.source)
        paths = export_normalized_data(tables, args.output)
        print(f"Exported {len(tables['dim_accounts'])} accounts to {paths['sqlite']}")
        if args.pdf:
            matches = find_accounts(tables["dim_accounts"], args.pdf)
            if matches.empty:
                raise SystemExit(f"No account found matching: {args.pdf}")
            pdf_path = create_printable_pdf(matches.iloc[0], tables["fact_payments"], tables["account_balance_verification"], args.output / "printables")
            print(f"Created printable: {pdf_path}")
        return
    run_ui()


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(f"ERROR: {error}", file=sys.stderr)
        raise
